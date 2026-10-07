"""RPA de custos: lê o portal do fornecedor e atualiza o custos.xlsx com segurança.

Fluxo:
    portal do fornecedor (SIMULADO)  ->  rpa_custos.py  ->  custos.xlsx  ->  auditor_margem.py

Regras:
    - custo igual ao atual ................ SEM MUDANÇA
    - variação até LIMITE_VARIACAO_PCT ..... ATUALIZADO (gravado no custos.xlsx)
    - variação acima do limite ............. CONFERIR (NÃO é gravado: pode ser erro do fornecedor)
    - custo ilegível no portal ............. ERRO DE LEITURA (o robô segue com os demais produtos)
    - SKU do portal que não está na planilha  SKU NOVO (não é adicionado, só listado)
    - SKU da planilha que não está no portal  FORA DO PORTAL (não é alterado, só listado)

Registros da execução:
    saida/logs/rpa_custos_AAAA-MM-DD.log ... histórico (horário, SKU, custo antes/depois, status)
    saida/screenshots/ ............... foto da tela quando o robô falha no navegador
    saida/relatorio_custos.xlsx ...... resumo + mudanças + SKUs novos + erros

Uso:
    python demo/portal_fornecedor/app.py          # terminal 1: portal simulado

    python src/rpa_custos.py --demo               # terminal 2: demonstração (não altera os originais)
    python src/rpa_custos.py --demo --visivel     # idem, mostrando o navegador
    python src/rpa_custos.py --simular            # compara e gera o relatório, sem gravar nada
    python src/rpa_custos.py                      # atualiza dados/custos.xlsx (com backup em saida/backups/)

Credenciais: nunca ficam no código nem no log. Vêm do .env (PORTAL_USUARIO / PORTAL_SENHA)
ou são digitadas no terminal. No portal simulado: demo / demo123 (fictícios).
"""
import argparse
import getpass
import logging
import os
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

load_dotenv()

# ================= CONFIGURAÇÃO =================
PORTAL_URL = os.getenv("PORTAL_URL", "http://127.0.0.1:5000")
PASTA_SESSAO = Path("dados/.playwright-session")          # login salvo (fica fora do Git)
ARQUIVO_SESSAO = PASTA_SESSAO / "estado.json"
ARQUIVO_PORTAL_CSV = "saida/custos_portal.csv"            # cópia do que foi lido no portal
TIMEOUT_MS = 10_000
MAX_PAGINAS = 200                                   # trava de segurança contra loop infinito

ARQUIVO_CUSTOS = "dados/custos.xlsx"
COL_SKU = "SKU"                                     # mesmos nomes usados pelo auditor
COL_CUSTO = "Custo do Fornecedor"
LIMITE_VARIACAO_PCT = 30.0                          # acima disso o custo NÃO é aplicado
PASTA_BACKUP = Path("saida/backups")
PASTA_LOGS = Path("saida/logs")
PASTA_SCREENSHOTS = Path("saida/screenshots")
ARQUIVO_RELATORIO = "saida/relatorio_custos.xlsx"
PASTA_DEMO = Path("demo")

# Todos os seletores num lugar só: se o portal mudar, é aqui que se ajusta.
SELETORES = {
    "campo_usuario": '[data-testid="campo-usuario"]',
    "campo_senha": '[data-testid="campo-senha"]',
    "botao_entrar": '[data-testid="botao-entrar"]',
    "erro_login": '[data-testid="erro-login"]',
    "tabela": '[data-testid="tabela-produtos"]',
    "linha": '[data-testid="linha-produto"]',
    "proxima": 'a[data-testid="proxima-pagina"]',  # só existe como <a> quando há próxima página
    "sku": '[data-campo="sku"]',
    "descricao": '[data-campo="descricao"]',
    "custo": '[data-campo="custo"]',
}
# ================================================

logger = logging.getLogger("rpa_custos")


class LoginError(Exception):
    pass


# ---------------------- log, falhas e screenshots ----------------------

def configurar_log():
    """Cria logs/rpa_custos_AAAA-MM-DD.log (várias execuções no dia vão para o mesmo arquivo)."""
    PASTA_LOGS.mkdir(parents=True, exist_ok=True)
    arquivo = PASTA_LOGS / f"rpa_custos_{datetime.now():%Y-%m-%d}.log"
    if not logger.handlers:
        handler = logging.FileHandler(arquivo, encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S"))
        logger.addHandler(handler)
        logger.setLevel(logging.INFO)
    return arquivo


def tirar_screenshot(page, motivo):
    """Foto da página inteira. Nunca deixa uma falha na foto esconder o erro original."""
    try:
        PASTA_SCREENSHOTS.mkdir(parents=True, exist_ok=True)
        caminho = PASTA_SCREENSHOTS / f"{datetime.now():%Y%m%d_%H%M%S}_{motivo}.png"
        page.screenshot(path=str(caminho), full_page=True)
        return caminho
    except Exception:
        return None


def falhar(mensagem, page=None, motivo="erro"):
    """Registra no log, tira screenshot (se houver página) e encerra com mensagem clara."""
    foto = tirar_screenshot(page, motivo) if page is not None else None
    logger.error("%s%s", mensagem, f" | screenshot={foto}" if foto else "")
    sys.exit(f"\n❌ {mensagem}" + (f"\n   Screenshot: {foto}" if foto else ""))


def num(valor):
    """Número para o log: 290.0 -> '290.00'; vazio -> '-'."""
    return "-" if valor is None or pd.isna(valor) else f"{valor:.2f}"


# ---------------------- valores e SKUs ----------------------

def normalizar_sku(valor):
    return str(valor).strip().upper()


def parse_brl(texto):
    """Converte texto monetário em float, sem confundir milhar e decimal.

    'R$ 1.234,56' -> 1234.56 | '849,90' -> 849.9 | '849.90' -> 849.9
    '1.234' -> 1234.0 (3 dígitos após um único ponto = milhar, convenção BR)
    '1,234.56' -> 1234.56 (formato americano: o último separador é o decimal)
    """
    s = re.sub(r"[^\d,.\-]", "", str(texto))  # tira "R$", espaços, etc.
    if not re.search(r"\d", s):
        raise ValueError(f"valor monetário inválido: {texto!r}")

    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):          # 1.234,56
            s = s.replace(".", "").replace(",", ".")
        else:                                    # 1,234.56
            s = s.replace(",", "")
    elif "," in s:                               # 849,90
        s = s.replace(",", ".")
    elif "." in s:
        depois = s.split(".")[-1]
        if s.count(".") > 1 or len(depois) == 3:  # 1.234 ou 1.234.567 -> milhar
            s = s.replace(".", "")
    return float(s)


def brl(valor):
    """1234.5 -> 'R$ 1.234,50' (para exibir no terminal)."""
    texto = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {texto}"


# ---------------------- portal (Playwright) ----------------------

def sessao_ativa(page):
    """Abre a tabela: se o portal redirecionar para /login, não há sessão válida."""
    page.goto(f"{PORTAL_URL}/produtos", wait_until="domcontentloaded")
    return "/login" not in page.url


def obter_credenciais():
    usuario = os.getenv("PORTAL_USUARIO") or input("Usuário do portal: ").strip()
    senha = os.getenv("PORTAL_SENHA") or getpass.getpass("Senha do portal: ")
    return usuario, senha


def fazer_login(page):
    usuario, senha = obter_credenciais()
    page.goto(f"{PORTAL_URL}/login", wait_until="domcontentloaded")
    page.fill(SELETORES["campo_usuario"], usuario)
    page.fill(SELETORES["campo_senha"], senha)
    page.click(SELETORES["botao_entrar"])
    try:
        # Espera a tabela (login ok) OU a mensagem de erro: assim a falha aparece na hora.
        page.wait_for_selector(f'{SELETORES["tabela"]}, {SELETORES["erro_login"]}', timeout=TIMEOUT_MS)
    except PlaywrightTimeout:
        raise LoginError("o portal não confirmou o login")
    erro = page.locator(SELETORES["erro_login"])
    if erro.count():
        raise LoginError(erro.inner_text().strip())


def ler_pagina(page, numero_pagina):
    """Lê as linhas da página atual. Uma linha com problema vira erro e o robô segue."""
    page.wait_for_selector(SELETORES["tabela"])
    registros, erros = [], []
    for indice, linha in enumerate(page.locator(SELETORES["linha"]).all(), start=1):
        sku = custo_texto = None
        try:
            sku = linha.locator(SELETORES["sku"]).inner_text().strip()
            custo_texto = linha.locator(SELETORES["custo"]).inner_text().strip()
            registros.append({
                "sku": sku,
                "descricao": linha.locator(SELETORES["descricao"]).inner_text().strip(),
                "custo_texto": custo_texto,
                "custo": parse_brl(custo_texto),
                "pagina": numero_pagina,
            })
        except (ValueError, PlaywrightTimeout) as e:
            motivo = str(e).splitlines()[0]
            erros.append({"pagina": numero_pagina, "linha": indice, "sku": sku,
                          "valor_lido": custo_texto, "erro": motivo})
            logger.warning("SKU=%s | pagina=%s | linha=%s | valor_lido=%r | STATUS=ERRO DE LEITURA | %s",
                           sku, numero_pagina, indice, custo_texto, motivo)
    return registros, erros


def extrair_todos(page):
    """Percorre as páginas clicando em 'Próxima' até ela deixar de ser um link."""
    page.goto(f"{PORTAL_URL}/produtos", wait_until="domcontentloaded")
    registros, erros, pagina = [], [], 1
    while True:
        lidos, falhas = ler_pagina(page, pagina)
        registros += lidos
        erros += falhas
        aviso = f" ({len(falhas)} com erro de leitura)" if falhas else ""
        print(f"  página {pagina}: {len(lidos)} produtos{aviso}")
        logger.info("Página %s lida: %s produtos, %s erros de leitura", pagina, len(lidos), len(falhas))

        if page.locator(SELETORES["proxima"]).count() == 0 or pagina >= MAX_PAGINAS:
            break
        page.click(SELETORES["proxima"])
        page.wait_for_url(re.compile(rf"pagina={pagina + 1}\b"))
        pagina += 1
    return registros, erros


def ler_portal(visivel):
    """Abre o navegador, garante o login (reaproveitando a sessão) e lê a tabela."""
    print(f"Portal: {PORTAL_URL}\n")
    with sync_playwright() as p:
        navegador = p.chromium.launch(headless=not visivel, slow_mo=300 if visivel else 0)
        contexto = navegador.new_context(
            storage_state=str(ARQUIVO_SESSAO) if ARQUIVO_SESSAO.exists() else None
        )
        page = contexto.new_page()
        page.set_default_timeout(TIMEOUT_MS)

        try:
            if sessao_ativa(page):
                print("Sessão salva reutilizada (sem novo login).")
                logger.info("Sessão salva reutilizada")
            else:
                print("Fazendo login...")
                fazer_login(page)
                PASTA_SESSAO.mkdir(parents=True, exist_ok=True)
                contexto.storage_state(path=str(ARQUIVO_SESSAO))
                print("Login realizado; sessão salva para as próximas execuções.")
                logger.info("Login realizado; sessão salva")

            print("Lendo a tabela de custos:")
            return extrair_todos(page)

        except LoginError as e:
            falhar(f"Falha no login: {e}", page, "login")
        except PlaywrightTimeout as e:
            falhar(f"Tempo esgotado esperando um elemento da página: {str(e).splitlines()[0]}", page, "timeout")
        except PlaywrightError as e:
            if "ERR_CONNECTION_REFUSED" in str(e):
                falhar(f"Não consegui abrir {PORTAL_URL}. O portal simulado está rodando?\n"
                       "   Suba com: python demo/portal_fornecedor/app.py")
            tirar_screenshot(page, "erro_playwright")
            raise
        finally:
            contexto.close()
            navegador.close()


# ---------------------- planilha de custos ----------------------

def carregar_planilha(caminho):
    """Abre o custos.xlsx preservando formatação. Devolve (workbook, aba, coluna_custo, itens).

    itens: {sku_normalizado: {"sku", "custo", "linhas": [números das linhas]}}
    O mesmo SKU pode aparecer em várias linhas (peça compatível com vários carros).
    """
    if not Path(caminho).is_file():
        falhar(f"Arquivo não encontrado: {caminho}")

    wb = load_workbook(caminho)
    ws = wb.worksheets[0]
    cabecalho = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}
    for nome in (COL_SKU, COL_CUSTO):
        if nome.lower() not in cabecalho:
            falhar(f"Coluna '{nome}' não encontrada em {caminho}. Colunas: {list(cabecalho)}")
    col_sku, col_custo = cabecalho[COL_SKU.lower()], cabecalho[COL_CUSTO.lower()]

    itens, divergentes = {}, set()
    for linha in range(2, ws.max_row + 1):
        sku = ws.cell(linha, col_sku).value
        if sku is None or str(sku).strip() == "":
            continue
        if isinstance(sku, float) and sku.is_integer():
            sku = int(sku)

        valor = ws.cell(linha, col_custo).value
        try:
            custo = None if valor is None else (float(valor) if isinstance(valor, (int, float)) else parse_brl(valor))
        except ValueError:
            custo = None

        chave = normalizar_sku(sku)
        item = itens.setdefault(chave, {"sku": str(sku).strip(), "custo": custo, "linhas": []})
        if item["linhas"] and custo != item["custo"]:
            divergentes.add(item["sku"])
        item["linhas"].append(linha)

    if divergentes:
        aviso = f"SKUs repetidos com custos diferentes na planilha: {', '.join(sorted(divergentes))}"
        print(f"⚠️  {aviso}\n    (todas as linhas do SKU receberão o mesmo custo novo)\n")
        logger.warning(aviso)
    return wb, ws, col_custo, itens


# ---------------------- regras ----------------------

def comparar(itens, df_portal, limite, erros):
    """Aplica as regras. Devolve (mudancas, skus_novos, fora_do_portal)."""
    portal = {normalizar_sku(r.sku): r for r in df_portal.itertuples()}
    com_erro = {normalizar_sku(e["sku"]): e for e in erros if e["sku"]}
    mudancas, fora = [], []

    for chave, item in itens.items():
        r = portal.get(chave)
        if r is None and chave in com_erro:  # estava no portal, mas o custo não pôde ser lido
            mudancas.append({
                "sku": item["sku"], "descricao": None, "custo_anterior": item["custo"],
                "custo_novo": None, "variacao_pct": None, "status": "ERRO DE LEITURA",
                "motivo": f"portal mostra {com_erro[chave]['valor_lido']!r}", "_chave": chave,
            })
            continue
        if r is None:
            fora.append({"sku": item["sku"], "custo_atual": item["custo"]})
            continue

        atual, novo = item["custo"], r.custo
        variacao, motivo = None, ""
        if not atual:  # vazio ou zero: não há base para comparar
            status, motivo = "CONFERIR", "sem custo anterior para comparar"
        else:
            variacao = (novo - atual) / atual * 100
            if abs(novo - atual) < 0.005:
                status = "SEM MUDANÇA"
            elif abs(variacao) > limite:
                status, motivo = "CONFERIR", f"variação acima de {limite:g}%"
            else:
                status = "ATUALIZADO"

        mudancas.append({
            "sku": item["sku"], "descricao": r.descricao,
            "custo_anterior": atual, "custo_novo": novo,
            "variacao_pct": None if variacao is None else round(variacao, 2),
            "status": status, "motivo": motivo, "_chave": chave,
        })

    novos = [{"sku": r.sku, "descricao": r.descricao, "custo_portal": r.custo}
             for chave, r in portal.items() if chave not in itens]
    return mudancas, novos, fora


def gravar_custos(wb, ws, col_custo, itens, mudancas, destino):
    """Escreve só as células de custo dos SKUs ATUALIZADO; o resto da planilha fica intacto."""
    for m in mudancas:
        if m["status"] == "ATUALIZADO":
            for linha in itens[m["_chave"]]["linhas"]:
                ws.cell(linha, col_custo).value = round(m["custo_novo"], 2)
    try:
        wb.save(destino)
    except PermissionError:
        falhar(f"Não consegui gravar {destino}. Feche o arquivo no Excel e rode de novo.")


# ---------------------- relatório ----------------------

CORES_STATUS = {"ATUALIZADO": "E2EFDA", "CONFERIR": "FFF2CC", "ERRO DE LEITURA": "F8CBAD"}


def formatar_aba(ws, colunas_status=None):
    """Cabeçalho em destaque, linha fixa, larguras ajustadas e status coloridos."""
    ws.freeze_panes = "A2"
    for celula in ws[1]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="D9E1F2")
    for coluna in ws.columns:
        maior = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
        ws.column_dimensions[coluna[0].column_letter].width = min(max(maior + 2, 10), 60)

    cabecalho = {c.value: c.column for c in ws[1]}
    if "status" in cabecalho:
        for linha in ws.iter_rows(min_row=2):
            cor = CORES_STATUS.get(linha[cabecalho["status"] - 1].value)
            if cor:
                for c in linha:
                    c.fill = PatternFill("solid", fgColor=cor)
    for nome, formato in (("custo_anterior", "#,##0.00"), ("custo_novo", "#,##0.00"),
                          ("custo_atual", "#,##0.00"), ("custo_portal", "#,##0.00"),
                          ("variacao_pct", "0.00")):
        if nome in cabecalho:
            for linha in ws.iter_rows(min_row=2, min_col=cabecalho[nome], max_col=cabecalho[nome]):
                linha[0].number_format = formato


def gerar_relatorio(mudancas, novos, fora, erros, resumo, caminho):
    ordem = {"ERRO DE LEITURA": 0, "CONFERIR": 1, "ATUALIZADO": 2, "SEM MUDANÇA": 3}
    df = pd.DataFrame(mudancas, columns=["sku", "descricao", "custo_anterior", "custo_novo",
                                         "variacao_pct", "status", "motivo", "_chave"]).drop(columns="_chave")
    df = df.sort_values("status", key=lambda s: s.map(ordem), kind="stable")

    abas = {
        "Resumo": pd.DataFrame(list(resumo.items()), columns=["campo", "valor"]),
        "Mudanças": df,
        "SKUs novos": pd.DataFrame(novos, columns=["sku", "descricao", "custo_portal"]),
        "Fora do portal": pd.DataFrame(fora, columns=["sku", "custo_atual"]),
        "Erros de leitura": pd.DataFrame(erros, columns=["pagina", "linha", "sku", "valor_lido", "erro"]),
    }

    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    try:
        with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
            for nome, tabela in abas.items():
                tabela.to_excel(xw, sheet_name=nome, index=False)
                formatar_aba(xw.sheets[nome])
    except PermissionError:
        falhar(f"Não consegui gravar {caminho}. Feche o arquivo no Excel e rode de novo.")


# ---------------------- execução ----------------------

def main():
    inicio = time.monotonic()
    parser = argparse.ArgumentParser(description="RPA de custos - portal do fornecedor -> custos.xlsx")
    parser.add_argument("--demo", action="store_true",
                        help=f"usa {PASTA_DEMO}/custos.xlsx e grava o resultado em uma cópia (originais intactos)")
    parser.add_argument("--simular", action="store_true", help="compara e gera o relatório, sem gravar o custos")
    parser.add_argument("--custos", default=None, help=f"planilha de custos (padrão: {ARQUIVO_CUSTOS})")
    parser.add_argument("--limite", type=float, default=LIMITE_VARIACAO_PCT,
                        help=f"variação máxima (%%) aplicada automaticamente (padrão: {LIMITE_VARIACAO_PCT:g})")
    parser.add_argument("--visivel", action="store_true", help="mostra o navegador (padrão: sem janela)")
    parser.add_argument("--limpar-sessao", action="store_true", help="apaga o login salvo antes de começar")
    args = parser.parse_args()

    origem = Path(args.custos or (PASTA_DEMO / "custos.xlsx" if args.demo else ARQUIVO_CUSTOS))
    relatorio = PASTA_DEMO / "saida" / "relatorio_custos_demo.xlsx" if args.demo else Path(ARQUIVO_RELATORIO)
    modo = "demo" if args.demo else ("simulação" if args.simular else "normal")

    arquivo_log = configurar_log()
    logger.info("INÍCIO | modo=%s | planilha=%s | portal=%s | limite=%g%%", modo, origem, PORTAL_URL, args.limite)

    if args.limpar_sessao and ARQUIVO_SESSAO.exists():
        ARQUIVO_SESSAO.unlink()
        print("Sessão salva removida.")
        logger.info("Sessão salva removida a pedido do usuário")

    print("=" * 40)
    print(" RPA - ATUALIZAÇÃO DE CUSTOS")
    print("=" * 40)
    print(f"Planilha de custos: {origem}")
    print(f"Limite automático : variação de até {args.limite:g}%\n")

    # Abre a planilha ANTES do navegador: se o arquivo estiver errado, falha rápido.
    wb, ws, col_custo, itens = carregar_planilha(origem)

    registros, erros = ler_portal(args.visivel)
    if not registros:
        falhar("Nenhum produto foi lido no portal. Nada foi alterado.")
    df_portal = pd.DataFrame(registros)
    arquivo_portal_csv = PASTA_DEMO / "saida" / "custos_portal_demo.csv" if args.demo else Path(ARQUIVO_PORTAL_CSV)
    arquivo_portal_csv.parent.mkdir(parents=True, exist_ok=True)
    df_portal.to_csv(arquivo_portal_csv, index=False, encoding="utf-8-sig")

    duplicados = df_portal[df_portal.duplicated("sku", keep=False)]["sku"].unique()
    if len(duplicados):
        aviso = f"SKUs repetidos no portal (vale o último): {', '.join(duplicados)}"
        print(f"\n⚠️  {aviso}")
        logger.warning(aviso)

    mudancas, novos, fora = comparar(itens, df_portal, args.limite, erros)
    for m in mudancas:
        var = "-" if m["variacao_pct"] is None else f"{m['variacao_pct']:+.1f}%"
        logger.info("SKU=%s | ANTES=%s | NOVO=%s | VAR=%s | STATUS=%s%s", m["sku"], num(m["custo_anterior"]),
                    num(m["custo_novo"]), var, m["status"], f" | {m['motivo']}" if m["motivo"] else "")
    contagem = {s: sum(1 for m in mudancas if m["status"] == s)
                for s in ("ATUALIZADO", "SEM MUDANÇA", "CONFERIR", "ERRO DE LEITURA")}

    # ----- gravação -----
    if args.simular:
        destino_msg = "Modo simulação: nenhum arquivo de custos foi alterado."
    elif contagem["ATUALIZADO"] == 0:
        destino_msg = "Nenhum custo a atualizar: planilha mantida como está."
    elif args.demo:
        destino = PASTA_DEMO / "saida" / "custos_atualizado.xlsx"  # o demo/custos.xlsx original não é tocado
        destino.parent.mkdir(parents=True, exist_ok=True)
        gravar_custos(wb, ws, col_custo, itens, mudancas, destino)
        destino_msg = f"Planilha atualizada gravada em: {destino}  (original intacto)"
    else:
        PASTA_BACKUP.mkdir(parents=True, exist_ok=True)
        backup = PASTA_BACKUP / f"{origem.stem}_{datetime.now():%Y%m%d_%H%M%S}{origem.suffix}"
        shutil.copy2(origem, backup)
        gravar_custos(wb, ws, col_custo, itens, mudancas, origem)
        destino_msg = f"{origem} atualizado. Backup do anterior em: {backup}"
    logger.info(destino_msg)

    duracao = int(time.monotonic() - inicio)
    tempo = f"{duracao // 60:02d}:{duracao % 60:02d}"
    resumo = {
        "Data e hora": f"{datetime.now():%d/%m/%Y %H:%M:%S}",
        "Modo": modo,
        "Tempo de execução": tempo,
        "Portal": PORTAL_URL,
        "Planilha de custos": str(origem),
        "Limite de variação automática (%)": args.limite,
        "Produtos lidos no portal": len(df_portal),
        "Atualizados": contagem["ATUALIZADO"],
        "Sem mudança": contagem["SEM MUDANÇA"],
        "Para conferir (não gravados)": contagem["CONFERIR"],
        "Erros de leitura": len(erros),
        "SKUs novos (não adicionados)": len(novos),
        "Na planilha, fora do portal": len(fora),
        "Log": str(arquivo_log),
    }
    gerar_relatorio(mudancas, novos, fora, erros, resumo, relatorio)
    logger.info("FIM | tempo=%s | atualizados=%s | sem_mudanca=%s | conferir=%s | erros_leitura=%s | relatorio=%s",
                tempo, contagem["ATUALIZADO"], contagem["SEM MUDANÇA"], contagem["CONFERIR"], len(erros), relatorio)

    # ----- resumo -----
    print("\n" + "=" * 40)
    print(" EXECUÇÃO FINALIZADA")
    print("=" * 40)
    print(f"Produtos lidos no portal    : {len(df_portal)}")
    print(f"Atualizados                 : {contagem['ATUALIZADO']}")
    print(f"Sem mudança                 : {contagem['SEM MUDANÇA']}")
    print(f"Para conferir (não gravados): {contagem['CONFERIR']}")
    print(f"Erros de leitura            : {len(erros)}")
    print(f"SKUs novos (não adicionados): {len(novos)}")
    print(f"Na planilha, fora do portal : {len(fora)}")
    print(f"Tempo de execução           : {tempo}")

    conferir = [m for m in mudancas if m["status"] == "CONFERIR"]
    if conferir:
        print("\n⚠️  CONFERIR com o fornecedor antes de aplicar:")
        for m in conferir:
            if m["variacao_pct"] is None:
                print(f"  {m['sku']:<10} {brl(m['custo_novo'])}  ({m['motivo']})")
            else:
                variacao = f"{m['variacao_pct']:+.1f}".replace(".", ",")
                print(f"  {m['sku']:<10} {brl(m['custo_anterior'])} -> {brl(m['custo_novo'])}  ({variacao}%)")
    if erros:
        print("\n❌ Não consegui ler o custo destes produtos (os demais foram processados):")
        for e in erros:
            print(f"  {e['sku'] or '(sem SKU)':<10} página {e['pagina']}, linha {e['linha']}: {e['valor_lido']!r}")

    print(f"\n{destino_msg}")
    print(f"Relatório: {relatorio}")
    print(f"Log      : {arquivo_log}")


if __name__ == "__main__":
    try:
        main()
    except Exception:  # erro inesperado: fica registrado no log com o traceback completo
        logger.exception("Erro inesperado")
        raise